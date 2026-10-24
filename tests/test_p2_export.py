"""Task 5: Excel report built on the sample workbook template."""
import datetime as dt

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import db, export, seed, web
from app.schemas import Record


@pytest.fixture(autouse=True)
def fresh_db(tmp_path, monkeypatch):
    monkeypatch.setattr(db, "DB_PATH", tmp_path / "test.db")
    db.init_db()
    seed.seed()


def _records() -> tuple[int, int]:
    base = dict(farmer_chat_id=42, farmer_name="María", lang="es", date="2026-10-21")
    a = db.insert_record(Record(**base, activity_key="treatment", parcel_alias="P5", product="Oxicloruro de cobre 50% WP",
                                target_pest_key="fungal_general", quantity=40, unit="kg", confidence=0.9,
                                transcript_original="Catorce kilos de cobre en los almendros, toda la parcela."))
    db.update_status(a, "farmer_confirmed", "farmer")
    db.set_judge(a, 0.41, "9.5 kg/ha copper is above the usual range")
    db.update_status(a, "needs_review")
    db.update_field(a, "quantity", 14, "Farmer said catorce (14)", "agronomist", reason_lang="en")
    db.update_status(a, "corrected", "agronomist")
    b = db.insert_record(Record(**base | {"lang": "en", "date": "2026-10-22"}, activity_key="irrigation",
                                parcel_alias="P2", quantity=120, unit="m3", area_ha=1.8,
                                transcript_original="Watered the Gala block today, about 120 cubic metres."))
    return a, b


def _rows(ws) -> list[dict]:
    header = [str(c.value).split(" (")[0] for c in ws[1]]
    return [dict(zip(header, row)) for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None]


def test_layout_matches_template(tmp_path):
    template = load_workbook(seed.SAMPLE_XLSX)
    out = load_workbook(export.build_report(str(tmp_path / "r.xlsx")))
    assert out.sheetnames == template.sheetnames
    for name in ("Parcels_DUN", "Field_Log", "Corrections", "Vocabulary"):
        assert [c.value for c in out[name][1]] == [c.value for c in template[name][1]], name


def test_empty_database_gives_empty_log(tmp_path):
    out = load_workbook(export.build_report(str(tmp_path / "r.xlsx")))
    assert _rows(out["Field_Log"]) == [] and _rows(out["Corrections"]) == []
    assert len(_rows(out["Parcels_DUN"])) == 5  # reference data is always there


def test_field_log_values_and_formulas(tmp_path):
    a, b = _records()
    ws = load_workbook(export.build_report(str(tmp_path / "r.xlsx")))["Field_Log"]
    rows = _rows(ws)
    assert [r["record_id"] for r in rows] == [f"R{a:04d}", f"R{b:04d}"]  # oldest first, sample rows gone
    first = rows[0]
    assert first["date"] == dt.datetime(2026, 10, 21) and first["parcel_alias"] == "P5"
    assert (first["quantity"], first["unit"], first["status"]) == (14, "kg", "corrected")
    assert first["area_ha"] == 4.2  # "toda la parcela": the parcel's area
    assert first["transcript_original"].startswith("Catorce")
    # Formulas point at their own row, lookups reach row 500
    assert ws["H3"].value.startswith('=IF(G3="","",') and "Parcels_DUN!$L$2:$L$500" in ws["H3"].value
    assert ws["M3"].value.startswith('=IF(L3="","",') and "Vocabulary!$D$2:$D$500" in ws["M3"].value
    assert ws["W3"].value == '=IF(AND(ISNUMBER(T3),ISNUMBER(V3)),IF(V3>0,T3/V3,""),"")'
    assert ws["F2"].value == "=E2"
    # Colours: blue values, green lookups, yellow review cells
    assert ws["A3"].font.color.rgb == "FF0000FF" and ws["H3"].font.color.rgb == "FF008000"
    assert ws["AD3"].fill.fgColor.rgb == "FFFFF2B3"
    assert ws.auto_filter.ref == "A1:AH3"


def test_corrections_tab(tmp_path):
    a, _ = _records()
    (c,) = _rows(load_workbook(export.build_report(str(tmp_path / "r.xlsx")))["Corrections"])
    assert (c["record_id"], c["field"], c["before"], c["after"]) == (f"R{a:04d}", "quantity", "40", "14")
    assert c["corrected_by"] == "agronomist" and c["used_as_fewshot"] == "yes"
    assert isinstance(c["timestamp"], dt.datetime) and c["timestamp"].tzinfo is None


def test_parcels_keep_formulas(tmp_path):
    ws = load_workbook(export.build_report(str(tmp_path / "r.xlsx")))["Parcels_DUN"]
    assert ws["A2"].value == "P1" and "первый участок" in ws["B2"].value
    assert ws["L2"].value == '=E2&"-"&F2&"-"&G2&"-"&H2&"-"&I2&"-"&J2&"-"&K2'
    assert "Vocabulary!$D$2:$D$500" in ws["O6"].value and "N6" in ws["O6"].value


def test_download_endpoint(monkeypatch):
    monkeypatch.delenv("WEB_USER", raising=False)
    monkeypatch.delenv("WEB_PASSWORD", raising=False)
    _records()
    with TestClient(web.app) as client:
        assert client.get("/api/report.xlsx").status_code == 401
        r = client.get("/api/report.xlsx", auth=("admin", "admin"))
    assert r.status_code == 200 and r.content[:2] == b"PK"  # xlsx is a zip
    assert "spreadsheetml" in r.headers["content-type"] and ".xlsx" in r.headers["content-disposition"]
