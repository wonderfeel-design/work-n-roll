"""Create the database tables and load the reference data from data/satech_dun_field_log_sample.xlsx:
Parcels_DUN, Vocabulary and UI_Strings. Idempotent (re-running updates the reference rows).

The workbook's Field_Log and Corrections rows are format examples and are never imported.

    python -m app.seed           # create tables, load / refresh reference data
    python -m app.seed --reset   # drop every table first (loses all records and corrections)
"""
import sys

from openpyxl import load_workbook
from psycopg import sql

from app import ROOT, db

SAMPLE_XLSX = ROOT / "data" / "satech_dun_field_log_sample.xlsx"
SIGPAC_PARTS = db.SIGPAC_PARTS
TABLES = ("corrections", "field_log", "farmers", "parcels", "ui_strings", "vocabulary")


def _rows(ws):
    header = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is not None:
            yield dict(zip(header, row))


def _int(v):
    return int(v) if v is not None else None


def _upsert(conn, table: str, row: dict, key: tuple[str, ...]) -> None:
    cols = list(row)
    conn.execute(
        sql.SQL("INSERT INTO {} ({}) VALUES ({}) ON CONFLICT ({}) DO UPDATE SET {}").format(
            sql.Identifier(table), sql.SQL(", ").join(map(sql.Identifier, cols)),
            sql.SQL(", ").join(sql.Placeholder() * len(cols)), sql.SQL(", ").join(map(sql.Identifier, key)),
            sql.SQL(", ").join(sql.SQL("{0} = excluded.{0}").format(sql.Identifier(c)) for c in cols if c not in key)),
        list(row.values()),
    )


def seed() -> None:
    wb = load_workbook(SAMPLE_XLSX)
    with db._connect() as conn:
        vocab = [{k: r[k] for k in ("key", "type", "en", "ca", "es", "ru")} for r in _rows(wb["Vocabulary"])]
        for v in vocab:
            _upsert(conn, "vocabulary", v, ("type", "key"))

        for r in _rows(wb["UI_Strings"]):
            _upsert(conn, "ui_strings", {k: r[k] for k in ("key", "en", "ca", "es", "ru")}, ("key",))

        crops = {v["key"] for v in vocab if v["type"] == "crop"}
        for r in _rows(wb["Parcels_DUN"]):
            if r["crop_key"] not in crops:
                raise ValueError(f"parcel {r['parcel_alias']}: crop_key {r['crop_key']!r} is not in Vocabulary")
            parts = {k: _int(r[k]) for k in SIGPAC_PARTS}
            _upsert(conn, "parcels", {
                "alias": r["parcel_alias"], "spoken_names": r["spoken_names"], "campaign": str(r["campaign"]),
                "explotacio_id": r["explotacio_id"], **parts,
                "sigpac_ref": "-".join(str(v) for v in parts.values()),
                **{k: r[k] for k in ("sigpac_use", "crop_key", "variety", "area_ha", "irrigation", "tenure",
                                     "eco_scheme")},
            }, ("alias",))


def reset() -> None:
    """Drop every table: the only way records and corrections ever disappear."""
    with db._connect() as conn:
        conn.execute(sql.SQL("DROP TABLE IF EXISTS {} CASCADE").format(sql.SQL(", ").join(map(sql.Identifier, TABLES))))


def main() -> None:
    if "--reset" in sys.argv:
        reset()
    db.init_db()
    seed()
    print(f"PostgreSQL schema {db.SCHEMA_NAME!r}: {len(db.get_parcels())} parcels, "
          f"{len(db.get_vocabulary())} vocabulary terms, {len(db.list_records())} records")


if __name__ == "__main__":
    main()
